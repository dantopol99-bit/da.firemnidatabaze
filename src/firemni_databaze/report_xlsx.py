"""Datová příloha reportu (XLSX): každá tabulka na vlastním listu, list „Metodika“
s ukazateli z katalogu a list „Zdroj a licence“.

Rozložení listu tabulky (spoléhá na něj report_kontrola):
    1: název tabulky   2: citace zdroje   3: označení odvozených údajů
    5: záhlaví (Položka | sloupce tabulky… | Poznámka)   6…: řádky tabulky
"""

from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

RADEK_ZAHLAVI = 5
TUCNE = Font(bold=True)


def _sirky(ws, sirky: list[int]) -> None:
    for i, s in enumerate(sirky, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = s


def zapis_xlsx(vysledek: dict, cesta: Path, katalog_cesta: Path) -> None:
    meta = vysledek["meta"]
    katalog = yaml.safe_load(katalog_cesta.read_text(encoding="utf-8"))["ukazatele"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Přehled"
    radky = [
        ("Oborově-regionální report – datová příloha", None),
        ("Obor", meta["obor_popis"]),
        ("Klasifikace", meta["klasifikace"]),
        ("Území", f"{meta['uzemi']['nazev']} ({meta['uzemi']['kod']})"),
        ("Datum snímku RES", meta["zadani"]["datum_snimku"]),
        ("Populace", meta["populace"]),
        ("Srovnávací kraje (zadal uživatel)", ", ".join(meta["zadani"]["srovnani"]) or "nezadány"),
        ("Práh zveřejnění (subjektů)", meta["pravidla"]["prah"]),
        ("Minimální rozsah reportu (subjektů)", meta["pravidla"]["min_rozsah"]),
        ("Práh spolehlivosti zařazení (%)", meta["pravidla"]["prah_spolehlivosti_pct"]),
        ("Citace", meta["citace"]),
        ("Označení", meta["oznaceni"]),
    ] + [("Varování", v) for v in meta["varovani"]] + [
        (None, None),
        ("Tabulky", None),
    ] + [(t["kod"], t["nazev"] + ("" if t["zverejneno"] else f" – nezveřejněno: {t['duvod']}"))
         for t in vysledek["tabulky"]]
    for r in radky:
        ws.append(list(r))
    ws["A1"].font = TUCNE
    _sirky(ws, [34, 120])

    for t in vysledek["tabulky"]:
        ws = wb.create_sheet(t["kod"][:31])
        ws.append([t["nazev"]])
        ws.append([meta["citace"]])
        ws.append([meta["oznaceni"]])
        ws.append([])
        ws["A1"].font = TUCNE
        if not t["zverejneno"]:
            ws.append([f"Nezveřejněno: {t['duvod']}"])
            continue
        ws.append(["Položka"] + [s["nazev"] for s in t["sloupce"]] + ["Poznámka"])
        for c in ws[RADEK_ZAHLAVI]:
            c.font = TUCNE
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in t["radky"]:
            ws.append([r["popis"]] + [r["hodnoty"].get(s["kod"]) for s in t["sloupce"]] + [r.get("poznamka")])
            if r["typ"] in ("celkem", "mezisoucet"):
                ws.cell(row=ws.max_row, column=1).font = TUCNE
        ws.append([])
        for p in t["poznamky"]:
            ws.append([p])
        _sirky(ws, [48] + [16] * len(t["sloupce"]) + [50])

    ws = wb.create_sheet("Metodika")
    pouzite = {r.get("ukazatele", {}).get(s["kod"]) or s["ukazatel"]
               for t in vysledek["tabulky"] for r in t["radky"] for s in t["sloupce"]}
    pole = ("kod", "nazev", "typ", "definice", "vzorec", "zdroj", "uroven", "chybejici_hodnoty", "vyklad", "stav")
    ws.append(["Kód", "Název", "Typ", "Definice", "Vzorec", "Zdroj", "Úroveň", "Ošetření chybějících hodnot",
               "Poznámka k výkladu", "Stav", "Použit v této příloze"])
    for c in ws[1]:
        c.font = TUCNE
    for u in katalog:
        ws.append([", ".join(u[p]) if isinstance(u.get(p), list) else u.get(p) for p in pole]
                  + ["ano" if u["kod"] in pouzite else "ne"])
    ws.append([])
    ws.append(["Pravidla ochrany malých buněk: buňky pod prahem se v každém členění slučují do „ostatní“; je-li "
               "„ostatní“ pod prahem, přidá se další nejmenší buňka; skryté číslo nesmí jít dopočítat ze součtů ani "
               "z jiné tabulky (ověřeno kontrolou dopočtu). Podrobnosti: docs/report_metodika.md."])
    _sirky(ws, [24, 40, 10, 60, 40, 40, 16, 60, 60, 14, 10])

    ws = wb.create_sheet("Zdroj a licence")
    for r in [
        ("Zdroj", meta["citace"]),
        ("Označení", meta["oznaceni"]),
        ("Licence", f"{meta['licence']['nazev']} – {meta['licence']['url']}"),
        ("Datum snímku RES", meta["zadani"]["datum_snimku"]),
        ("Soubor RES", "https://opendata.csu.gov.cz/soubory/od/od_org03/res_data.csv"),
    ] + [("Další zdroj", z) for z in meta["dalsi_zdroje"]] + [
        ("Povinnost", "Upravené nebo odvozené údaje musí být označeny jako odvozené a nesmí být prezentovány "
                      "jako nezměněné oficiální statistiky ČSÚ."),
        ("Vygenerováno", meta["vygenerovano"]),
    ]:
        ws.append(list(r))
    _sirky(ws, [22, 140])
    wb.save(cesta)
