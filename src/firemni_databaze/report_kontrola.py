"""Kontrola výstupu reportu před zveřejněním.

Spuštění:
    python -m firemni_databaze.report_kontrola reporty/vystupy/F__CZ051__2026-09-15
    python -m firemni_databaze.report_kontrola reporty/vystupy/*

Ověřuje:
  1. každé číslo ve vysledek.json je ve stejné buňce priloha.xlsx (a list tabulky existuje),
  2. každé číslo má ukazatel z katalogu (reporty/katalog_ukazatelu.yaml),
  3. žádný zveřejněný počet (typ „pocet“) není pod prahem,
  4. odvozené ukazatele (podíl, index, pořadí, průměr) jsou jen u zveřejněných počtů,
  5. skryté číslo nejde dopočítat odečtem od součtu ani z jiné tabulky
     (lineární soustava z _interni/kontrola.json), „ostatní“ má aspoň 2 položky,
  6. čísla ve výstupu odpovídají interním proměnným,
  7. citace zdroje (ČSÚ, RES, datum snímku, CC BY 4.0) a označení odvozených údajů
     jsou v JSON i na každém listu XLSX,
  8. texty neobsahují „firmy“ ani „aktivní“ (rozhodnutí 1); výjimkou je jednotka ČSÚ „aktivní podnik“.
Interní soubor _interni/kontrola.json obsahuje i skrytá čísla – nezveřejňuje se.
"""

import argparse
import json
import re
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

from firemni_databaze.report import KATALOG, OZNACENI
from firemni_databaze.report_potlaceni import dopocitatelne
from firemni_databaze.report_xlsx import RADEK_ZAHLAVI

ZAKAZANA_SLOVA = re.compile(r"\bfirm\w*|\baktivn\w*+(?!\s+podnik)", re.IGNORECASE)
# „aktivní podnik(y)“ je oficiální jednotka ČSÚ v demografii podniků (RESDP00, rozhodnutí 10),
# nikoli označení registrovaných subjektů – proto je povolené jen v tomto sousloví.
ODVOZENE_TYPY = ("podil", "index", "poradi", "prumer")


def _texty(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if k not in ("kod", "ukazatel", "ukazatele", "promenna", "zaklad", "zaklady", "url"):
                yield from _texty(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _texty(v)


def zkontroluj(adresar: Path, katalog_cesta: Path = KATALOG) -> list[str]:
    """Vrátí seznam chyb (prázdný = výstup je v pořádku)."""
    chyby: list[str] = []
    vysledek = json.loads((adresar / "vysledek.json").read_text(encoding="utf-8"))
    interni = json.loads((adresar / "_interni" / "kontrola.json").read_text(encoding="utf-8"))
    katalog = {u["kod"]: u for u in yaml.safe_load(katalog_cesta.read_text(encoding="utf-8"))["ukazatele"]}
    wb = load_workbook(adresar / "priloha.xlsx", read_only=False)
    meta = vysledek["meta"]
    prah = meta["zadani"]["prah"]
    promenne = interni["promenne"]

    # 7. citace a označení
    datum = meta["zadani"]["datum_snimku"]
    r, m, dd = datum.split("-")
    for nutne in ("Český statistický úřad", "Registr ekonomických subjektů", "CC BY 4.0",
                  f"{int(dd)}. {int(m)}. {r}"):
        if nutne not in meta["citace"]:
            chyby.append(f"citace neobsahuje „{nutne}“")
    if meta.get("oznaceni") != OZNACENI:
        chyby.append("chybí označení odvozených údajů")
    for list_ in ("Metodika", "Zdroj a licence"):
        if list_ not in wb.sheetnames:
            chyby.append(f"v XLSX chybí list „{list_}“")

    for t in vysledek["tabulky"]:
        kod = t["kod"]
        if kod[:31] not in wb.sheetnames:
            chyby.append(f"{kod}: v XLSX chybí list")
            continue
        ws = wb[kod[:31]]
        if ws.cell(2, 1).value != meta["citace"] or ws.cell(3, 1).value != meta["oznaceni"]:
            chyby.append(f"{kod}: list nemá citaci a označení")
        if not t["zverejneno"]:
            if t["radky"]:
                chyby.append(f"{kod}: nezveřejněná tabulka má řádky")
            continue
        for i, radek in enumerate(t["radky"]):
            xr = RADEK_ZAHLAVI + 1 + i
            if ws.cell(xr, 1).value != radek["popis"]:
                chyby.append(f"{kod}: řádek {i} popis v XLSX {ws.cell(xr, 1).value!r} ≠ {radek['popis']!r}")
            for j, sl in enumerate(t["sloupce"]):
                # základ (počet), z něhož je číslo odvozené; tabulky s více územími ho mají po sloupcích
                zaklad = radek.get("zaklady", {}).get(sl["kod"]) or radek.get("zaklad") or radek.get("promenna")
                hodnota = radek["hodnoty"].get(sl["kod"])
                if hodnota is None:
                    continue
                misto = f"{kod}/{radek['popis']}/{sl['kod']}"
                # 1. JSON = XLSX
                x = ws.cell(xr, 2 + j).value
                if x != hodnota:
                    chyby.append(f"{misto}: XLSX {x!r} ≠ JSON {hodnota!r}")
                # 2. katalog
                uk = radek.get("ukazatele", {}).get(sl["kod"]) or sl["ukazatel"]
                if uk not in katalog:
                    chyby.append(f"{misto}: ukazatel {uk!r} není v katalogu")
                    continue
                typ = katalog[uk]["typ"]
                # 3. práh
                if typ == "pocet" and hodnota < prah:
                    chyby.append(f"{misto}: zveřejněný počet {hodnota} je pod prahem {prah}")
                # 4. odvozené jen u zveřejněných počtů
                if typ in ODVOZENE_TYPY and zaklad and not promenne.get(zaklad, {}).get("zverejneno"):
                    chyby.append(f"{misto}: odvozený ukazatel u skrytého počtu ({zaklad})")
                # 6. počet odpovídá interní proměnné
                var = radek.get("zaklady", {}).get(sl["kod"]) or (
                    radek.get("promenna") if sl["kod"] in ("pocet", "hodnota") else None)
                if typ == "pocet" and var:
                    v = promenne.get(var)
                    if v is None or v["hodnota"] != hodnota or not v["zverejneno"]:
                        chyby.append(f"{misto}: neodpovídá interní proměnné {var}")

    # 5. dopočet
    zverejnene = {k for k, v in promenne.items() if v["zverejneno"]}
    skryte = {k for k, v in promenne.items() if not v["zverejneno"] and v.get("chranit", True)}
    rovnice = [r if isinstance(r, dict) else tuple(r) for r in interni["rovnice"]]
    for v in dopocitatelne(rovnice, zverejnene, skryte):
        chyby.append(f"skryté číslo {v} jde dopočítat ze zveřejněných")
    for celek, casti in (r for r in rovnice if isinstance(r, tuple)):
        if celek.startswith("ostatni:"):
            if len(casti) < 2:
                chyby.append(f"{celek}: „ostatní“ má jen {len(casti)} položku")
            if promenne[celek]["zverejneno"] and promenne[celek]["hodnota"] < prah:
                chyby.append(f"{celek}: „ostatní“ je pod prahem")

    # 8. zakázaná slova (JSON texty i všechny textové buňky XLSX)
    for text in list(_texty(vysledek)) + [c for ws in wb for row in ws.iter_rows(values_only=True)
                                          for c in row if isinstance(c, str)]:
        if ZAKAZANA_SLOVA.search(text) and "AKTIVNI_PODIL" not in text:
            chyby.append(f"zakázané slovo v textu: {text[:80]!r}")
    return chyby


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("adresare", nargs="+", type=Path)
    args = parser.parse_args()
    vse_ok = True
    for a in args.adresare:
        if not (a / "vysledek.json").exists():
            continue
        chyby = zkontroluj(a)
        print(f"{a.name}: {'OK' if not chyby else f'{len(chyby)} chyb'}")
        for c in chyby:
            print(f"  - {c}")
        vse_ok &= not chyby
    return 0 if vse_ok else 1


if __name__ == "__main__":
    sys.exit(main())


# ---------------------------------------------------------------------------
# Konzistence PDF (Blok 3): každé číslo v textu PDF musí být v JSON i v XLSX
# ---------------------------------------------------------------------------

# Číslo v sazbě: tisíce oddělené úzkou (U+202F) nebo pevnou mezerou, desetinná čárka.
_CISLO_SAZBA = re.compile(r"\d+(?:[  ]\d{3})*(?:[,.]\d+)?")
# V textech JSON/XLSX se tisíce oddělují i obyčejnou mezerou („na 1 000 obyvatel“).
_CISLO_TEXT = re.compile(r"\d{1,3}(?: \d{3})+(?:[,.]\d+)?")
_STRANA = re.compile(r"strana\s+\d+\s+z\s+\d+", re.IGNORECASE)


def _norm(cislo: str) -> str:
    from decimal import Decimal, InvalidOperation
    t = re.sub(r"[\s  ]", "", cislo).replace(",", ".")
    try:
        return format(Decimal(t).normalize(), "f")
    except InvalidOperation:
        return t


def cisla_v_textu(text: str) -> set[str]:
    """Normalizovaná čísla v textu sazby (bez čísel stránek)."""
    return {_norm(c) for c in _CISLO_SAZBA.findall(_STRANA.sub(" ", text))}


def _cisla_zdroje(hodnoty) -> set[str]:
    vysledek: set[str] = set()
    for h in hodnoty:
        if isinstance(h, bool) or h is None:
            continue
        if isinstance(h, (int, float)):
            # znaménko se v textu sazby nečte (−5,8 i 5,8 je „5.8“), proto i absolutní hodnota
            vysledek |= {_norm(str(h)), _norm(str(abs(h)))}
        elif isinstance(h, str):
            vysledek |= {_norm(c) for c in _CISLO_SAZBA.findall(h)}
            vysledek |= {_norm(c) for c in _CISLO_TEXT.findall(h)}
    return vysledek


def _hodnoty_json(obj):
    if isinstance(obj, dict):
        for v in obj.values():
            yield from _hodnoty_json(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _hodnoty_json(v)
    else:
        yield obj


def povolena_cisla(adresar: Path) -> set[str]:
    """Čísla, která smí být v textu reportu: vyskytují se v JSON I v XLSX."""
    vysledek = json.loads((adresar / "vysledek.json").read_text(encoding="utf-8"))
    wb = load_workbook(adresar / "priloha.xlsx", read_only=True)
    v_xlsx = _cisla_zdroje(c for ws in wb for row in ws.iter_rows(values_only=True) for c in row)
    return _cisla_zdroje(_hodnoty_json(vysledek)) & v_xlsx


def text_pdf(pdf: Path) -> str:
    from pypdf import PdfReader
    return "\n".join(stranka.extract_text() or "" for stranka in PdfReader(pdf).pages)


def zkontroluj_pdf(pdf: Path, adresar: Path) -> list[str]:
    """Čísla v textu PDF, která nejsou v JSON a XLSX (prázdný seznam = v pořádku)."""
    povolena = povolena_cisla(adresar)
    text = text_pdf(pdf)
    chyby = []
    for c in sorted(cisla_v_textu(text) - povolena):
        misto = next((radek.strip() for radek in text.splitlines() if c.replace(".", ",") in radek.replace(" ", "")
                      or c in radek), "")
        chyby.append(f"číslo {c} v PDF není v JSON a XLSX (…{misto[:70]}…)")
    return chyby
