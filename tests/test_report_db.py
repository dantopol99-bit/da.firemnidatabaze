"""Výpočet reportu a kontrola výstupu nad databází se snímkem RES.

Bez databáze se schématem res a načteným snímkem se testy přeskočí.
Spuštění:  python -m unittest tests.test_report_db
"""

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import yaml

from firemni_databaze.report import FO_FORMY, KATALOG, MalyRozsah, Zadani, spocitej, uloz
from firemni_databaze.report_kontrola import zkontroluj
from firemni_databaze.report_xlsx import RADEK_ZAHLAVI
from tests.test_res_import_db import CONN, SNIMKY


def tabulka(vysledek, kod):
    return next(t for t in vysledek["tabulky"] if t["kod"] == kod)


def radek_t01(vysledek, ukazatel):
    return next(r for r in tabulka(vysledek, "T01_zakladni")["radky"] if r["ukazatele"]["hodnota"] == ukazatel)


@unittest.skipIf(CONN is None or not SNIMKY, "databáze se snímkem RES není dostupná")
class TestReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="report_test_"))
        cls.vystupy = {}
        for klic, zadani in {
            "F_LBK": Zadani(["F"], "CZ051", srovnani=["CZ052", "CZ041"]),
            "62_JES": Zadani(["62"], "Jeseník", min_rozsah=0),
            "41_42_CR": Zadani(["41", "42"], "CZ"),
        }.items():
            vysledek, interni = spocitej(CONN, zadani)
            cls.vystupy[klic] = (vysledek, interni, uloz(vysledek, interni, cls.tmp))
        # území pod prahem: třída s 1–9 subjekty v okrese
        with CONN.cursor() as cur:
            cur.execute(
                "SELECT n.trida, s.okreslau FROM res.subjekt s JOIN res.cis_nace n "
                "ON n.klasifikace = 80004 AND n.kod = s.nace WHERE s.ddatzan IS NULL AND n.trida IS NOT NULL "
                "AND s.datum_snimku = res.posledni_snimek() GROUP BY 1, 2 HAVING count(*) BETWEEN 2 AND 9 LIMIT 1")
            trida, okres = cur.fetchone()
        CONN.rollback()
        vysledek, interni = spocitej(CONN, Zadani([trida], okres, min_rozsah=0))
        cls.vystupy["MALE"] = (vysledek, interni, uloz(vysledek, interni, cls.tmp))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)
        CONN.rollback()

    def test_kontrola_vystupu_projde(self):
        for klic, (_, _, adresar) in self.vystupy.items():
            self.assertEqual(zkontroluj(adresar), [], klic)

    def test_pocet_odpovida_databazi(self):
        vysledek, _, _ = self.vystupy["F_LBK"]
        with CONN.cursor() as cur:
            cur.execute("SELECT count(*) FROM res.v_subjekt WHERE datum_snimku = res.posledni_snimek() "
                        "AND NOT je_zanikly AND kraj_kod = 'CZ051' AND nace_sekce = 'F'")
            ocekavano = cur.fetchone()[0]
        CONN.rollback()
        self.assertEqual(radek_t01(vysledek, "REG_POCET")["hodnoty"]["hodnota"], ocekavano)

    def test_uzemi_pod_prahem_nic_nezverejni(self):
        vysledek, interni, _ = self.vystupy["MALE"]
        self.assertIsNone(radek_t01(vysledek, "REG_POCET")["hodnoty"]["hodnota"])
        for kod in ("T05_fo_po", "T06_pravni_forma", "T07_velikost_fo", "T08_velikost_po", "T09_vekova_struktura"):
            self.assertFalse(tabulka(vysledek, kod)["zverejneno"], kod)
        self.assertIsNone(radek_t01(vysledek, "VEK_PRUMER")["hodnoty"]["hodnota"])
        self.assertIsNone(radek_t01(vysledek, "HUSTOTA")["hodnoty"]["hodnota"])

    def test_rozdeleni_fo_po_pod_prahem(self):
        vysledek, interni, _ = self.vystupy["62_JES"]
        po = interni["promenne"]["fopo:PO"]["hodnota"]
        if po >= 10:
            self.skipTest("v Jeseníku už není PO pod prahem")
        self.assertFalse(tabulka(vysledek, "T05_fo_po")["zverejneno"])
        self.assertFalse(tabulka(vysledek, "T07_velikost_fo")["zverejneno"])
        self.assertFalse(tabulka(vysledek, "T08_velikost_po")["zverejneno"])

    def test_zadne_zverejnene_cislo_pod_prahem_v_xlsx(self):
        typy = {u["kod"]: u["typ"] for u in yaml.safe_load(KATALOG.read_text(encoding="utf-8"))["ukazatele"]}
        for klic, (vysledek, _, adresar) in self.vystupy.items():
            wb = load_workbook(adresar / "priloha.xlsx")
            for t in vysledek["tabulky"]:
                if not t["zverejneno"]:
                    continue
                for i, r in enumerate(t["radky"]):
                    for j, s in enumerate(t["sloupce"]):
                        uk = r.get("ukazatele", {}).get(s["kod"]) or s["ukazatel"]
                        if typy.get(uk) == "pocet":
                            v = wb[t["kod"]].cell(RADEK_ZAHLAVI + 1 + i, 2 + j).value
                            if isinstance(v, (int, float)):
                                self.assertGreaterEqual(v, 10, f"{klic}/{t['kod']}/{r['popis']}")

    # --- kontrola musí odhalit podvržené chyby -------------------------------------------
    def kopie(self, klic):
        cil = self.tmp / f"kopie_{klic}_{self._testMethodName}"
        shutil.copytree(self.vystupy[klic][2], cil)
        return cil

    def test_kontrola_odhali_rozdil_json_xlsx(self):
        a = self.kopie("F_LBK")
        wb = load_workbook(a / "priloha.xlsx")
        wb["T03_kraje"].cell(RADEK_ZAHLAVI + 1, 3).value = 1
        wb.save(a / "priloha.xlsx")
        self.assertTrue(any("XLSX" in c for c in zkontroluj(a)))

    def test_kontrola_odhali_cislo_pod_prahem(self):
        a = self.kopie("F_LBK")
        v = json.loads((a / "vysledek.json").read_text(encoding="utf-8"))
        t = tabulka(v, "T06_pravni_forma")
        t["radky"][0]["hodnoty"]["pocet"] = 7
        (a / "vysledek.json").write_text(json.dumps(v, ensure_ascii=False), encoding="utf-8")
        self.assertTrue(any("pod prahem" in c for c in zkontroluj(a)))

    def test_kontrola_odhali_dopocet(self):
        a = self.kopie("62_JES")
        i = json.loads((a / "_interni" / "kontrola.json").read_text(encoding="utf-8"))
        # zveřejni všechny členy nějakého „ostatní“ kromě jednoho → ten poslední jde dopočítat
        celek, casti = next((r[0], r[1]) for r in i["rovnice"]
                            if isinstance(r, list) and r[0].startswith("ostatni:") and len(r[1]) >= 2)
        for p in casti[:-1]:
            i["promenne"][p]["zverejneno"] = True
        (a / "_interni" / "kontrola.json").write_text(json.dumps(i), encoding="utf-8")
        self.assertTrue(any(f"skryté číslo {casti[-1]}" in c for c in zkontroluj(a)))

    def test_kontrola_odhali_zakazane_slovo(self):
        a = self.kopie("F_LBK")
        v = json.loads((a / "vysledek.json").read_text(encoding="utf-8"))
        v["tabulky"][0]["poznamky"].append("Počet aktivních firem v kraji.")
        (a / "vysledek.json").write_text(json.dumps(v, ensure_ascii=False), encoding="utf-8")
        self.assertTrue(any("zakázané slovo" in c for c in zkontroluj(a)))

    def test_minimalni_rozsah_odmitne_s_navrhem(self):
        with self.assertRaises(MalyRozsah) as ctx:
            spocitej(CONN, Zadani(["62"], "Jeseník"))
        CONN.rollback()
        e = ctx.exception
        self.assertLess(e.pocet, 100)
        popisy = [n["popis"] for n in e.navrhy]
        self.assertTrue(any("Olomoucký kraj" in p and p.startswith("62") for p in popisy), popisy)
        self.assertTrue(any(p.startswith("J ") for p in popisy), popisy)
        self.assertTrue(all(isinstance(n["pocet"], int) for n in e.navrhy))
        self.assertIn("Návrh vyšší úrovně", e.vypis())

    def test_spolehlivost_zarazeni(self):
        vysledek, _, _ = self.vystupy["62_JES"]
        r = radek_t01(vysledek, "SPOLEHLIVOST_JEN_SEKCE")
        self.assertGreater(r["hodnoty"]["hodnota"], 25)
        self.assertTrue(vysledek["meta"]["varovani"])
        self.assertEqual(self.vystupy["F_LBK"][0]["meta"]["varovani"], [])

    # --- Blok 4: benchmarky, míra zániku PO, zjištění, čeština -----------------------------
    def test_benchmark_cr_odpovida_databazi(self):
        vysledek, _, _ = self.vystupy["F_LBK"]
        with CONN.cursor() as cur:
            cur.execute("SELECT round(100.0 * count(*) FILTER (WHERE forma = ANY(%s)) / count(*), 1) FROM res.v_subjekt "
                        "WHERE datum_snimku = res.posledni_snimek() AND NOT je_zanikly AND nace_sekce = 'F'",
                        (sorted(FO_FORMY),))
            podil_fo_cr = float(cur.fetchone()[0])
        CONN.rollback()
        fo = next(r for r in tabulka(vysledek, "T13_bench_fo_po")["radky"] if r["promenna"] == "fopo:FO")
        self.assertEqual(fo["hodnoty"]["cr"], podil_fo_cr)
        self.assertEqual(fo["hodnoty"]["uzemi"], next(
            r for r in tabulka(vysledek, "T05_fo_po")["radky"] if r["promenna"] == "fopo:FO")["hodnoty"]["podil"])
        self.assertIn("kraj_CZ052", fo["hodnoty"])

    def test_benchmark_jen_nad_prahem(self):
        """Každý zveřejněný podíl v benchmarku stojí na zveřejněném počtu aspoň na prahu."""
        for klic, (vysledek, interni, _) in self.vystupy.items():
            for t in vysledek["tabulky"]:
                if not t["kod"].startswith(("T13", "T14", "T15", "T16", "T17")) or not t["zverejneno"]:
                    continue
                for r in t["radky"]:
                    for sl, h in r["hodnoty"].items():
                        if h is None or sl == "rozdil_cr":
                            continue
                        p = interni["promenne"][r["zaklady"][sl]]
                        self.assertTrue(p["zverejneno"], f"{klic}/{t['kod']}/{r['popis']}/{sl}")
                        self.assertGreaterEqual(p["hodnota"], 10, f"{klic}/{t['kod']}/{r['popis']}/{sl}")

    def test_bez_benchmarku_u_cr_a_skryteho_uzemi(self):
        cr, _, _ = self.vystupy["41_42_CR"]
        self.assertFalse(any(t["kod"].startswith("T13") for t in cr["tabulky"]))
        male, _, _ = self.vystupy["MALE"]
        for t in male["tabulky"]:
            if t["kod"].startswith(("T13", "T14", "T15", "T16", "T17")):
                self.assertFalse(t["zverejneno"], t["kod"])

    def test_mira_zaniku_po(self):
        vysledek, interni, _ = self.vystupy["F_LBK"]
        t18 = tabulka(vysledek, "T18_mira_zaniku_po")
        t10 = {r["popis"][:4]: r["hodnoty"]["pocet"] for r in tabulka(vysledek, "T10_zaniky_po")["radky"]
               if r["typ"] == "polozka"}
        self.assertEqual({r["uzemi"] for r in t18["radky"]}, {"CZ051", "CZ", "CZ052", "CZ041"})
        for r in t18["radky"]:
            h = r["hodnoty"]
            if h["mira"] is not None:
                self.assertEqual(h["mira"], round(100 * h["zan"] / h["stav"], 2))
            if h["mira_obd"] is not None:
                self.assertEqual(h["mira_obd"], round(100 * h["obd"] / h["stav"], 2))
            if r["uzemi"] == "CZ051" and h["zan"] is not None:
                self.assertEqual(h["zan"], t10[str(r["rok"])])
            if r["rok"] == 2026:
                self.assertIsNone(h["mira"])      # neúplný rok: jen srovnatelné období
        # nezávislá rekonstrukce stavu PO v ČR k 1. 1. 2024 přímo v SQL
        with CONN.cursor() as cur:
            cur.execute("SELECT count(*) FROM res.subjekt s JOIN res.cis_nace n ON n.klasifikace = 80004 "
                        "AND n.kod = s.nace WHERE s.datum_snimku = res.posledni_snimek() AND n.sekce = 'F' "
                        "AND NOT n.je_pseudokod AND NOT (coalesce(s.forma, '') = ANY(%s)) "
                        "AND s.ddatvzn < '2024-01-01' AND (s.ddatzan IS NULL OR s.ddatzan >= '2024-01-01') "
                        "AND (s.ddatzan IS NULL OR s.forma IS NOT NULL)", (sorted(FO_FORMY),))
            stav = cur.fetchone()[0]
        CONN.rollback()
        cr2024 = next(r for r in t18["radky"] if r["uzemi"] == "CZ" and r["rok"] == 2024)
        self.assertEqual(cr2024["hodnoty"]["stav"], stav)
        self.assertTrue(any("Omezení" in p and "4 roky" in p for p in t18["poznamky"]))

    def test_zjisteni(self):
        vysledek, _, adresar = self.vystupy["F_LBK"]
        z = vysledek["zjisteni"]
        self.assertTrue(z)
        self.assertEqual([x["sila"] for x in z], sorted((x["sila"] for x in z), reverse=True))
        self.assertTrue(all(x["sila"] >= 0.03 for x in z))
        typy = {x["typ"] for x in z}
        for typ in ("odchylka_od_cr", "zmena_trendu", "rozdily_uvnitr_uzemi", "aktivita_oboru", "divergence_poradi"):
            self.assertIn(typ, typy)
        tab = {t["kod"]: t for t in vysledek["tabulky"]}
        for x in z:
            for c in x["cisla"]:
                if c["tabulka"] == "T19_zjisteni":
                    continue
                r = next(r for r in tab[c["tabulka"]]["radky"] if r["popis"] == c["radek"])
                self.assertEqual(r["hodnoty"][c["sloupec"]], c["hodnota"], (x["id"], c))
        soubor = json.loads((adresar / "zjisteni.json").read_text(encoding="utf-8"))
        self.assertEqual(soubor["zjisteni"], z)

    def test_zadny_vystup_nema_v_uzemi_nazev(self):
        from firemni_databaze.report_cestina import tabulka as lokativy
        nazvy = [z["nazev"] for z in lokativy().values()]
        for klic, (_, _, adresar) in self.vystupy.items():
            texty = [(adresar / f).read_text(encoding="utf-8") for f in ("vysledek.json", "vysledek.md", "zjisteni.json")]
            wb = load_workbook(adresar / "priloha.xlsx")
            texty += [c for ws in wb for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
            for text in texty:
                for nazev in nazvy:
                    self.assertNotIn(f"v území {nazev}".lower(), text.lower(), klic)

    def test_neplatne_zadani(self):
        for zadani in (Zadani(["41", "4120"], "CZ"), Zadani(["00"], "CZ"), Zadani(["F"], "Neexistující"),
                       Zadani(["F"], "CZ051", srovnani=["Liberec"]), Zadani(["41201"], "CZ")):
            with self.assertRaises(ValueError, msg=str(zadani)):
                spocitej(CONN, zadani)
            CONN.rollback()


if __name__ == "__main__":
    unittest.main()
